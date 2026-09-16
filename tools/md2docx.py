# -*- coding: utf-8 -*-
"""Markdown -> .docx для учебной работы.

Оформление: Times New Roman 14 пт, интервал 1,5, выравнивание по ширине,
абзацный отступ 1,25 см, поля 2 см со всех сторон, таблицы с сеткой (11 пт),
нумерация страниц внизу по центру (на титульном листе номер не ставится).

Использование:
    python3 tools/md2docx.py ВХОД.md ВЫХОД.docx [--title ТИТУЛ.txt]

Формат файла титульного листа (обычный текст, одна строка — один абзац):
    обычная строка   -> по центру, 14 пт
    ^строка          -> по центру, 16 пт, полужирный
    >строка          -> по правому краю
    ~                -> пустая строка

Маркер <!--TOC--> в markdown заменяется на поле автособираемого оглавления.
"""
import re, sys, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Times New Roman'


def setup(doc):
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for i, sz in ((1, 16), (2, 15), (3, 14), (4, 14)):
        h = doc.styles[f'Heading {i}']
        h.font.name = FONT
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.paragraph_format.space_before = Pt(18 if i < 3 else 12)
        h.paragraph_format.space_after = Pt(12 if i < 3 else 6)
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.alignment = (WD_ALIGN_PARAGRAPH.CENTER if i == 1
                                        else WD_ALIGN_PARAGRAPH.LEFT)


def field(par, instr):
    """Вставить поле Word (PAGE, TOC и т. п.)."""
    r = par.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    for el in (b, t, sep, e):
        r._r.append(el)


def page_numbers(doc):
    """Номер страницы внизу по центру; на первой странице не выводится."""
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    field(p, 'PAGE')
    for r in p.runs:
        r.font.name = FONT
        r.font.size = Pt(12)
    sec.first_page_footer.paragraphs[0].text = ''


def title_page(doc, path):
    for raw in pathlib.Path(path).read_text(encoding='utf-8').rstrip('\n').split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        if raw.strip() == '~':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if raw.startswith('^'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(raw[1:]); r.bold = True; r.font.size = Pt(16)
        elif raw.startswith('>'):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(raw[1:]); r.font.size = Pt(14)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(raw); r.font.size = Pt(14)
        r.font.name = FONT
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    field(p, 'TOC \\o "2-3" \\h \\z \\u')
    note = doc.add_paragraph()
    note.paragraph_format.first_line_indent = Cm(0)
    r = note.add_run('(Оглавление собирается автоматически: щёлкните по нему '
                     'и нажмите F9, затем «Обновить целиком».)')
    r.italic = True; r.font.size = Pt(11); r.font.name = FONT


INLINE = re.compile(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*|`.+?`)', re.S)


def add_runs(par, text, size=None):
    text = re.sub(r'\\([*_\\])', r'\1', text)
    for part in INLINE.split(text):
        if not part:
            continue
        b = i = mono = False
        if part.startswith('***') and part.endswith('***') and len(part) > 6:
            part, b, i = part[3:-3], True, True
        elif part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, i = part[1:-1], True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            part, mono = part[1:-1], True
        r = par.add_run(part)
        r.bold, r.italic = b, i
        r.font.name = 'Courier New' if mono else FONT
        r.font.size = Pt(11) if mono else (Pt(size) if size else None)


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def add_table(doc, rows):
    head, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(head))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell in enumerate(head):
        p = t.rows[0].cells[j].paragraphs[0]
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, f'**{cell}**', size=11)
    for i, row in enumerate(body, start=1):
        for j in range(len(head)):
            p = t.rows[i].cells[j].paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, row[j] if j < len(row) else '', size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def convert(src, dst, titul=None):
    lines = pathlib.Path(src).read_text(encoding='utf-8').split('\n')
    doc = Document()
    setup(doc)
    page_numbers(doc)
    if titul:
        title_page(doc, titul)

    i, n = 0, len(lines)
    while i < n:
        ln, s = lines[i], lines[i].strip()

        if s == '<!--TOC-->':
            toc(doc); i += 1; continue

        if s == '<!--PAGEBREAK-->':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); i += 1; continue

        if s.startswith('```'):
            i += 1; buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Cm(1)
                r = p.add_run(b); r.font.name = 'Courier New'; r.font.size = Pt(10)
            doc.add_paragraph()
            continue

        if s.startswith('|') and i + 1 < n and re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            rows = [split_row(s)]; i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i])); i += 1
            add_table(doc, rows)
            continue

        if s.startswith('#'):
            lvl = len(s) - len(s.lstrip('#'))
            doc.add_heading('', level=min(lvl, 4))
            add_runs(doc.paragraphs[-1], s.lstrip('# ').strip())
            i += 1; continue

        if re.match(r'^(---+|\*\*\*+)$', s):
            i += 1; continue

        if s.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.first_line_indent = Cm(0)
            add_runs(p, s.lstrip('> ').strip(), size=12)
            i += 1; continue

        m = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2))
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            i += 1; continue

        if re.match(r'^[-*]\s+', s):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^[-*]\s+', '', s))
            p.paragraph_format.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            i += 1; continue

        if not s:
            i += 1; continue

        add_runs(doc.add_paragraph(), s)
        i += 1

    doc.save(dst)
    print('OK ->', dst)


if __name__ == '__main__':
    args = sys.argv[1:]
    titul = None
    if '--title' in args:
        k = args.index('--title'); titul = args[k + 1]; args = args[:k] + args[k + 2:]
    convert(args[0], args[1], titul)
